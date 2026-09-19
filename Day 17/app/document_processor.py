"""
Document processing: validates uploads and extracts raw text from
PDF, DOCX, and TXT files.

Every failure mode raises DocumentProcessingError with an http_status
and a human-readable message, so main.py can turn it straight into an
HTTPException without a chain of if/elif in the route handler.
"""
import io
import os
from dataclasses import dataclass

from pypdf import PdfReader
from pypdf.errors import PdfReadError
import docx as docx_lib

from app.config import ALLOWED_EXTENSIONS, MAX_FILE_SIZE_BYTES


class DocumentProcessingError(Exception):
    def __init__(self, http_status: int, message: str):
        self.http_status = http_status
        self.message = message
        super().__init__(message)


@dataclass
class ExtractedDocument:
    text: str
    num_pages: int | None = None


def validate_upload(filename: str, content_type: str, raw_bytes: bytes) -> str:
    """Runs all pre-extraction checks. Returns the validated file extension."""
    if not filename:
        raise DocumentProcessingError(400, "Filename is missing.")

    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentProcessingError(
            415,
            f"Unsupported file type '{ext or 'unknown'}'. "
            f"Allowed types: {', '.join(sorted(ALLOWED_EXTENSIONS))}.",
        )

    if len(raw_bytes) == 0:
        raise DocumentProcessingError(400, "Uploaded file is empty (0 bytes).")

    if len(raw_bytes) > MAX_FILE_SIZE_BYTES:
        raise DocumentProcessingError(
            413,
            f"File too large ({len(raw_bytes)} bytes). "
            f"Maximum allowed is {MAX_FILE_SIZE_BYTES} bytes.",
        )

    return ext


def extract_text(ext: str, raw_bytes: bytes) -> ExtractedDocument:
    """Dispatches to the right extractor and normalizes failures."""
    try:
        if ext == ".pdf":
            return _extract_pdf(raw_bytes)
        elif ext == ".docx":
            return _extract_docx(raw_bytes)
        elif ext == ".txt":
            return _extract_txt(raw_bytes)
        else:
            # Should be unreachable given validate_upload, but stay defensive.
            raise DocumentProcessingError(415, f"Unsupported file type '{ext}'.")
    except DocumentProcessingError:
        raise
    except (PdfReadError, ValueError, KeyError) as e:
        raise DocumentProcessingError(
            400, f"File appears to be corrupted or unreadable: {e}"
        )
    except Exception as e:  # last-resort guard so a bad file never 500s silently
        raise DocumentProcessingError(
            400, f"Could not process file due to an unexpected error: {e}"
        )


def _extract_pdf(raw_bytes: bytes) -> ExtractedDocument:
    reader = PdfReader(io.BytesIO(raw_bytes))
    if reader.is_encrypted:
        try:
            reader.decrypt("")  # try an empty password before giving up
        except Exception:
            pass
    if reader.is_encrypted:
        raise DocumentProcessingError(400, "PDF is password-protected and cannot be read.")

    pages_text = []
    for page in reader.pages:
        pages_text.append(page.extract_text() or "")
    text = "\n".join(pages_text).strip()

    if not text:
        raise DocumentProcessingError(
            422,
            "No extractable text found in PDF. It may be a scanned/image-only "
            "document that requires OCR, which this API does not perform.",
        )
    return ExtractedDocument(text=text, num_pages=len(reader.pages))


def _extract_docx(raw_bytes: bytes) -> ExtractedDocument:
    doc = docx_lib.Document(io.BytesIO(raw_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs).strip()
    if not text:
        raise DocumentProcessingError(422, "No extractable text found in DOCX file.")
    return ExtractedDocument(text=text)


def _extract_txt(raw_bytes: bytes) -> ExtractedDocument:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            text = raw_bytes.decode(encoding).strip()
            if not text:
                raise DocumentProcessingError(422, "Text file is empty after decoding.")
            return ExtractedDocument(text=text)
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise DocumentProcessingError(400, "Could not decode text file with any supported encoding.")
